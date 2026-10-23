"""Modèle économique ScanLinkPay — classeur Excel (toutes les valeurs calculées par formules)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.comments import Comment

OUT = 'ScanLinkPay_Modele_Economique.xlsx'  # puis l'ouvrir dans Excel ou LibreOffice : les formules se calculent
MONTHS = 36
FIRST = 3  # column C = mois 1

NAVY = '0B1F6B'
BRAND = '1A3BE0'
LIGHT = 'EEF2FF'
F = 'Arial'
font = lambda **k: Font(name=F, size=k.pop('size', 10), **k)
INPUT = font(color='0000FF')
LINK = font(color='008000')
BOLD = font(bold=True)
WHITE_B = font(bold=True, color='FFFFFF')
YELLOW = PatternFill('solid', fgColor='FFFF00')
HEAD = PatternFill('solid', fgColor=NAVY)
SUB = PatternFill('solid', fgColor=LIGHT)
TOTAL = PatternFill('solid', fgColor='DDE3FF')
thin = Side(style='thin', color='BFC7E8')
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)

USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
CDF = '#,##0" FC";(#,##0" FC");"-"'
GBP = '£#,##0;(£#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
PCT2 = '0.00%;(0.00%);"-"'
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'

wb = Workbook()
wb.remove(wb.active)


def sheet(name, widths):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[col(i)].width = w
    return ws


def title(ws, text, sub=None, width=8):
    ws['A1'] = text
    ws['A1'].font = font(size=16, bold=True, color=NAVY)
    if sub:
        ws['A2'] = sub
        ws['A2'].font = font(size=10, italic=True, color='555555')
    ws.row_dimensions[1].height = 24


def header(ws, row, labels, start=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=start + i, value=t)
        c.font = WHITE_B
        c.fill = HEAD
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BOX


def section(ws, row, text, span=6):
    c = ws.cell(row=row, column=1, value=text)
    c.font = font(bold=True, color='FFFFFF', size=11)
    for i in range(1, span + 1):
        ws.cell(row=row, column=i).fill = PatternFill('solid', fgColor=BRAND)


def put(ws, ref, value, fmt=None, f=None, fill=None, bold=False):
    c = ws[ref]
    c.value = value
    c.font = f or (BOLD if bold else font())
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


# =====================================================================================================================
# 1. Hypothèses (all prices, fees and costs — blue = to change)
# =====================================================================================================================
H = sheet('Hypothèses', [52, 16, 12, 70])
title(H, 'Hypothèses — tout ce qui se modifie', 'Bleu = valeur à discuter / modifier · Jaune = hypothèse clé · Les autres onglets se recalculent seuls.')
A = {}  # key -> absolute reference 'Hypothèses'!$B$n
r = 4


def hyp(key, label, value, fmt, note, key_hyp=False):
    global r
    put(H, f'A{r}', label)
    c = put(H, f'B{r}', value, fmt, INPUT, YELLOW if key_hyp else None)
    c.border = BOX
    put(H, f'D{r}', note, f=font(size=9, color='555555'))
    H[f'D{r}'].alignment = Alignment(wrap_text=True, vertical='top')
    A[key] = f"'Hypothèses'!$B${r}"
    r += 1


def hsec(text):
    global r
    r += 1
    section(H, r, text, 4)
    r += 1
    header(H, r, ['Paramètre', 'Valeur', '', 'Source / commentaire'])
    r += 1


hsec('Général')
put(H, f'A{r}', 'Scénario étudié')
c = put(H, f'B{r}', 'Réaliste', None, INPUT, YELLOW)
c.border = BOX
put(H, f'D{r}', 'Choisir Pessimiste, Réaliste ou Optimiste (liste). Les paramètres de chaque scénario sont dans l’onglet Scénarios.', f=font(size=9, color='555555'))
dv = DataValidation(type='list', formula1='"Pessimiste,Réaliste,Optimiste"', allow_blank=False)
H.add_data_validation(dv)
dv.add(f'B{r}')
A['scenario'] = f"'Hypothèses'!$B${r}"
r += 1
hyp('fx', 'Taux de change (CDF pour 1 USD)', 2850, NUM, 'Taux indicatif à mettre à jour (le super admin le règle aussi dans l’application). Le modèle est en USD.')
hyp('gbp', 'Taux de change (USD pour 1 GBP)', 1.27, '0.00', 'Pour l’objectif exprimé en livres sterling (onglet Rentabilité). À mettre à jour.')
hyp('vat', 'TVA RDC sur nos services', 0.16, PCT, 'TVA standard RDC 16 %. Mettre 0 si non assujetti. À confirmer avec le comptable.')
hyp('vat_on', 'Nos prix incluent la TVA ? (1 = oui, 0 = non)', 1, '0', '1 = prudent : une part de chaque recette est reversée à l’État.', True)
hyp('ibp', 'Impôt sur les bénéfices (IBP)', 0.30, PCT, 'Taux de droit commun RDC (30 %). À confirmer avec le comptable.')

hsec('Paiements chez les commerçants et entreprises (tarifs du code)')
hyp('comm', 'Commission ScanLinkPay sur un paiement', 0.025, PCT2, 'Règle par défaut du code (USD et CDF : 2,5 %, min 0,05 $, max 100 $). Modifiable par le super admin → Commissions.', True)
hyp('payin_cost', 'Frais FlexPaie à l’encaissement supportés par ScanLinkPay', 0.0, PCT2, 'FlexPaie facture ~2,5 % (3 % par carte) au PAYEUR, en plus du montant (docs/FLEXPAIE.md). Coût 0 pour ScanLinkPay tant que c’est le cas.')
hyp('wd_share', 'Part de l’argent reçu retirée vers Mobile Money', 0.60, PCT, 'Hypothèse : le reste reste dans les portefeuilles ou repart en paiements.')
hyp('wd_fee', 'Frais de retrait facturés à l’utilisateur', 0.01, PCT2, 'Valeur actuelle du code : 1 %. ⚠ Inférieure au coût FlexPaie : voir l’onglet Retraits (recommandé ≥ 2,5 %).', True)
hyp('payout_cost', 'Coût FlexPaie d’un retrait (payout)', 0.02, PCT2, 'Environ 2 % selon FlexPaie (docs/FLEXPAIE.md). À confirmer sur le contrat.')
hyp('tr_fee', 'Frais de transfert entre portefeuilles', 0.0, PCT2, 'Gratuit dans le code (attirer les utilisateurs).')
hyp('tr_volume', 'Transferts par client actif et par mois ($)', 30, USD, 'Hypothèse.')
hyp('avg_pay', 'Montant moyen d’un paiement chez un commerçant ($)', 20, USD, 'Sert à compter les transactions (onglet Rentabilité).')

hsec('Abonnements (prix réglés par le super admin)')
hyp('sub_price', 'Abonnement entreprise — prix par mois ($)', 15, USD2, 'À remplacer par votre prix réel (super admin → Tarifs → Abonnements). Tout inclus : caisse, ventes, stock, rapports…', True)
hyp('sub_trial', 'Abonnement entreprise — essai gratuit (jours)', 30, NUM, 'Réglage super admin (informatif dans ce modèle).')
hyp('sup_price', 'Abonnement Fournisseur — prix par mois ($)', 10, USD2, 'Décision du 23/10 : 10 $ par mois (migration 084).', True)
hyp('sup_free', 'Abonnement Fournisseur — mois offerts sur l’année', 2, NUM, 'L’année = 12 mois pour le prix de (12 − N).')
hyp('sup_annual', 'Part des fournisseurs qui paient à l’année', 0.30, PCT, 'Hypothèse.')
hyp('pro_cdf', 'Mes dépenses « Pro » — prix par mois (CDF)', 5000, CDF, 'Valeur par défaut du code (5 000 CDF), réglable par le super admin.')

hsec('Applications partenaires qui utilisent l’API ScanLinkPay')
put(H, f'A{r}', 'Commission sur les paiements reçus par API : voir le tableau ci-dessous (une ligne par application).', f=font(size=9, italic=True))
r += 1
header(H, r, ['Application', 'Mois de début', 'Transactions / mois au début', 'Montant moyen ($)', 'Croissance mensuelle', 'Commission'], 1)
H.column_dimensions['C'].width = 14
H.column_dimensions['E'].width = 14
H.column_dimensions['F'].width = 12
r += 1
API_FIRST = r
apps = [
    ('HotelSaver (abonnements des hôtels + acomptes des clients)', 1, 40, 20, 0.06, 'Intégration en place : quick-pay, commission selon les règles du super admin.'),
    ('Application partenaire 2', 6, 100, 15, 0.08, 'Exemple : boutique en ligne, école, transport…'),
    ('Application partenaire 3', 12, 150, 10, 0.08, 'Exemple.'),
    ('Application partenaire 4', 18, 200, 10, 0.08, 'Exemple.'),
]
for name, start, tx, amt, g, note in apps:
    put(H, f'A{r}', name, f=INPUT)
    for c_, v, fm in (('B', start, NUM), ('C', tx, NUM), ('D', amt, USD), ('E', g, PCT)):
        cell = put(H, f'{c_}{r}', v, fm, INPUT)
        cell.border = BOX
    put(H, f'F{r}', f"={A['comm']}", PCT2, LINK).border = BOX
    H[f'A{r}'].comment = Comment(note, 'Modèle')
    r += 1
API_LAST = r - 1
put(H, f'A{r}', 'Mois de début 0 ou vide = application pas encore branchée. La commission suit par défaut la règle standard (modifiable ligne par ligne).', f=font(size=9, italic=True, color='555555'))
r += 1

hsec('Coûts fixes mensuels ($)')
fixed = [
    ('c_render', 'Hébergement API + site (Render)', 32, 'Render : API Standard ~25 $ + site ~7 $ (à vérifier sur la facture).'),
    ('c_supabase', 'Base de données (Supabase Pro)', 25, 'Projet production. « LinkPay Test » sur l’offre gratuite.'),
    ('c_redis', 'Redis (limites de requêtes)', 10, 'Si utilisé (REDIS_URL).'),
    ('c_domain', 'Domaines, e-mails, petits services', 10, ''),
    ('c_team', 'Équipe (développeurs)', 1000, 'Hypothèse : 2 développeurs × 500 $. Mettre 0 pour voir le résultat « avant salaires ».'),
    ('c_support', 'Support client et agents terrain', 300, 'Hypothèse.'),
    ('c_mkt', 'Marketing fixe (réseaux sociaux, affiches)', 300, 'Hypothèse.'),
    ('c_legal', 'Comptabilité, juridique, licences', 150, 'Hypothèse (agréments financiers, RCCM…).'),
    ('c_office', 'Bureau, internet, électricité', 150, 'Hypothèse.'),
]
for k, lbl, v, note in fixed:
    hyp(k, lbl, v, USD, note)
put(H, f'A{r}', 'Total coûts fixes mensuels', bold=True)
put(H, f'B{r}', '=SUM(' + ','.join(A[k].split('!')[1] for k, *_ in fixed) + ')', USD, bold=True, fill=TOTAL)
A['fixed_total'] = f"'Hypothèses'!$B${r}"
r += 1
hyp('team_growth', 'Hausse annuelle des coûts fixes (an 2 et an 3)', 0.25, PCT, 'Embauches, serveurs plus gros. Appliquée au début de chaque année.')

hsec('Coûts variables')
hyp('infra_1000', 'Serveurs : coût pour 1 000 utilisateurs actifs ($/mois)', 4, USD2, 'Au-delà de l’offre de base.')
hyp('sms_price', 'Prix d’un SMS ($)', 0.03, USD2, 'Alertes de sécurité, codes. À confirmer avec le fournisseur SMS.')
hyp('sms_new', 'SMS par nouvel utilisateur', 3, NUM, 'Inscription, vérification.')
hyp('cac_e', 'Coût pour gagner une entreprise ($)', 15, USD, 'Démonstration, déplacement, formation.')
hyp('cac_m', 'Coût pour gagner un commerçant ($)', 3, USD, 'Affiche QR, visite.')
hyp('cac_c', 'Coût pour gagner un client ($)', 0.3, USD2, 'Bouche-à-oreille, petites promotions.')

hsec('Démarrage (une seule fois)')
hyp('init_play', 'Compte Google Play (25 $)', 25, USD, 'Frais unique Google.')
hyp('init_legal', 'Création, enregistrement, agréments', 500, USD, 'Hypothèse.')
hyp('init_cards', 'Impression des premières cartes ScanLinkPay', 200, USD, 'Hypothèse.')
hyp('init_other', 'Matériel et divers', 300, USD, 'Hypothèse (téléphones de test, imprimante…).')
put(H, f'A{r}', 'Total investissement de départ', bold=True)
put(H, f'B{r}', f"=SUM({A['init_play'].split('!')[1]}:{A['init_other'].split('!')[1]})", USD, bold=True, fill=TOTAL)
A['init_total'] = f"'Hypothèses'!$B${r}"
r += 1

hsec('Objectif')
hyp('target_gbp', 'Bénéfice net mensuel visé (£)', 1000, GBP, 'Objectif cité le 22/10 : gagner 1 000 £. Voir l’onglet Rentabilité.', True)
H.freeze_panes = 'A4'

# =====================================================================================================================
# 2. Scénarios
# =====================================================================================================================
S = sheet('Scénarios', [50, 15, 15, 15, 15, 55])
title(S, 'Scénarios — rythme d’adoption', 'Colonne « Actif » = le scénario choisi dans Hypothèses. Modifier les colonnes bleues.')
header(S, 4, ['Paramètre', 'Pessimiste', 'Réaliste', 'Optimiste', 'Actif', 'Commentaire'])
SC = {}
scen = [
    ('e0', 'Entreprises au départ', 3, 5, 10, NUM, 'Entreprises déjà inscrites (supermarchés, électronique).'),
    ('m0', 'Commerçants indépendants au départ', 5, 10, 25, NUM, ''),
    ('c0', 'Clients actifs au départ', 100, 200, 500, NUM, ''),
    ('e_new', 'Nouvelles entreprises par mois (au début)', 1, 3, 6, NUM, ''),
    ('m_new', 'Nouveaux commerçants par mois (au début)', 5, 10, 20, NUM, ''),
    ('c_new', 'Nouveaux clients par mois (au début)', 80, 200, 400, NUM, ''),
    ('g1', 'Croissance mensuelle des inscriptions — an 1', 0.04, 0.08, 0.12, PCT, 'Les nouvelles inscriptions augmentent chaque mois de ce pourcentage.'),
    ('g23', 'Croissance mensuelle des inscriptions — ans 2 et 3', 0.02, 0.04, 0.06, PCT, ''),
    ('ch_em', 'Départs mensuels (entreprises, commerçants)', 0.05, 0.03, 0.02, PCT, 'Churn.'),
    ('ch_c', 'Départs mensuels (clients)', 0.08, 0.05, 0.04, PCT, ''),
    ('conv', 'Entreprises qui paient après l’essai', 0.40, 0.60, 0.75, PCT, ''),
    ('vol_m', 'Paiements reçus par commerçant et par mois ($)', 100, 300, 600, USD, 'Via QR / carte ScanLinkPay.'),
    ('vol_e', 'Paiements ScanLinkPay reçus par entreprise et par mois ($)', 300, 800, 1500, USD, 'Caisse + ventes : seulement la part payée par ScanLinkPay (QR, carte), pas les espèces — faible au début.'),
    ('pro', 'Clients qui prennent « Mes dépenses Pro »', 0.01, 0.02, 0.04, PCT, ''),
    ('sup_share', 'Entreprises qui deviennent fournisseurs', 0.08, 0.15, 0.25, PCT, ''),
    ('sup_conv', 'Fournisseurs qui paient après l’essai', 0.40, 0.60, 0.75, PCT, ''),
    ('ord', 'Commandes reçues par fournisseur et par mois', 4, 8, 15, NUM, ''),
    ('basket', 'Montant moyen d’une commande entre entreprises ($)', 150, 300, 500, USD, ''),
    ('ord_slp', 'Part des commandes payées par ScanLinkPay', 0.25, 0.40, 0.60, PCT, 'Le reste est payé en espèces / Mobile Money hors application (pas de commission).'),
    ('api_mult', 'Multiplicateur des volumes des applications partenaires', 0.5, 1.0, 1.5, '0.0x', 'Appliqué au tableau des applications (Hypothèses).'),
]
for i, (k, lbl, p, re_, o, fm, note) in enumerate(scen):
    row = 5 + i
    put(S, f'A{row}', lbl)
    for c_, v in (('B', p), ('C', re_), ('D', o)):
        put(S, f'{c_}{row}', v, fm, INPUT).border = BOX
    put(S, f'E{row}', f"=INDEX(B{row}:D{row},MATCH({A['scenario']},$B$4:$D$4,0))", fm, bold=True, fill=TOTAL).border = BOX
    put(S, f'F{row}', note, f=font(size=9, color='555555'))
    SC[k] = f"'Scénarios'!$E${row}"
S.freeze_panes = 'B5'

# =====================================================================================================================
# 3. Prévisions — 36 months
# =====================================================================================================================
P = sheet('Prévisions', [46, 8] + [12] * MONTHS)
title(P, 'Prévisions mensuelles sur 3 ans (USD)', 'Tout est calculé à partir des onglets Hypothèses et Scénarios. Ne rien saisir ici.')
LASTC = col(FIRST + MONTHS - 1)
put(P, 'A4', 'Mois', bold=True)
put(P, 'A5', 'Année', bold=True)
for m in range(1, MONTHS + 1):
    c = col(FIRST + m - 1)
    put(P, f'{c}4', m, NUM, WHITE_B, HEAD).alignment = Alignment(horizontal='center')
    put(P, f'{c}5', f'=INT(({c}4-1)/12)+1', '"An "0', bold=True, fill=SUB).alignment = Alignment(horizontal='center')
R = {}
prow = 6


def line(key, label, unit, formula_fn, fmt, style=None):
    """formula_fn(c, prev) -> formula for month column c (prev = previous column letter or None)."""
    global prow
    put(P, f'A{prow}', label, f=BOLD if style in ('total', 'key') else font())
    put(P, f'B{prow}', unit, f=font(size=9, color='777777'))
    for m in range(1, MONTHS + 1):
        c = col(FIRST + m - 1)
        prev = col(FIRST + m - 2) if m > 1 else None
        cell = put(P, f'{c}{prow}', formula_fn(c, prev, m), fmt)
        if style == 'total':
            cell.fill = TOTAL
            cell.font = BOLD
        if style == 'key':
            cell.font = BOLD
    if style == 'total':
        P[f'A{prow}'].fill = TOTAL
        P[f'B{prow}'].fill = TOTAL
    R[key] = prow
    prow += 1


def psec(text):
    global prow
    prow += 1
    section(P, prow, text, 2)
    for m in range(MONTHS):
        P.cell(row=prow, column=FIRST + m).fill = PatternFill('solid', fgColor=BRAND)
    prow += 1


growth = lambda c: f"IF({c}$5=1,{SC['g1']},{SC['g23']})"

psec('Utilisateurs')
line('g', 'Facteur de croissance des inscriptions', 'x', lambda c, p, m: '=1' if m == 1 else f'={p}{prow}*(1+{growth(c)})', '0.00')
G = R['g']
line('e_new', 'Nouvelles entreprises', 'nb', lambda c, p, m: f"={SC['e_new']}*{c}{G}", NUM1)
line('e', 'Entreprises actives', 'nb', lambda c, p, m: (f"={SC['e0']}*(1-{SC['ch_em']})+{c}{R['e_new']}" if m == 1 else f"={p}{prow}*(1-{SC['ch_em']})+{c}{R['e_new']}"), NUM1, 'key')
line('e_pay', 'Entreprises qui paient l’abonnement', 'nb', lambda c, p, m: f"={c}{R['e']}*{SC['conv']}", NUM1)
line('m_new', 'Nouveaux commerçants', 'nb', lambda c, p, m: f"={SC['m_new']}*{c}{G}", NUM1)
line('mer', 'Commerçants actifs', 'nb', lambda c, p, m: (f"={SC['m0']}*(1-{SC['ch_em']})+{c}{R['m_new']}" if m == 1 else f"={p}{prow}*(1-{SC['ch_em']})+{c}{R['m_new']}"), NUM1, 'key')
line('c_new', 'Nouveaux clients', 'nb', lambda c, p, m: f"={SC['c_new']}*{c}{G}", NUM1)
line('cli', 'Clients actifs', 'nb', lambda c, p, m: (f"={SC['c0']}*(1-{SC['ch_c']})+{c}{R['c_new']}" if m == 1 else f"={p}{prow}*(1-{SC['ch_c']})+{c}{R['c_new']}"), NUM, 'key')
line('sup', 'Fournisseurs (réseau entre entreprises)', 'nb', lambda c, p, m: f"={c}{R['e']}*{SC['sup_share']}", NUM1)
line('sup_pay', 'Fournisseurs qui paient l’abonnement Fournisseur', 'nb', lambda c, p, m: f"={c}{R['sup']}*{SC['sup_conv']}", NUM1)
line('users', 'Utilisateurs actifs (total)', 'nb', lambda c, p, m: f"={c}{R['e']}+{c}{R['mer']}+{c}{R['cli']}", NUM, 'total')

psec('Volumes d’argent')
line('vol_pay', 'Paiements reçus par commerçants et entreprises', '$', lambda c, p, m: f"={c}{R['mer']}*{SC['vol_m']}+{c}{R['e']}*{SC['vol_e']}", USD)
line('vol_b2b', 'Commandes entre entreprises payées par ScanLinkPay', '$', lambda c, p, m: f"={c}{R['sup_pay']}*{SC['ord']}*{SC['basket']}*{SC['ord_slp']}", USD)
for i in range(API_FIRST, API_LAST + 1):
    j = i - API_FIRST + 1
    line(f'api{j}', f"=\"API — \"&'Hypothèses'!$A${i}", '$',
         lambda c, p, m, i=i: (f"=IF(AND('Hypothèses'!$B${i}>0,{c}$4>='Hypothèses'!$B${i}),"
                               f"'Hypothèses'!$C${i}*(1+'Hypothèses'!$E${i})^({c}$4-'Hypothèses'!$B${i})*'Hypothèses'!$D${i}*{SC['api_mult']},0)"), USD)
line('vol_api', 'Paiements reçus via l’API (toutes applications)', '$', lambda c, p, m: f"=SUM({c}{R['api1']}:{c}{R['api' + str(API_LAST - API_FIRST + 1)]})", USD, 'key')
line('vol_wd', 'Retraits vers Mobile Money', '$', lambda c, p, m: f"=({c}{R['vol_pay']}+{c}{R['vol_b2b']})*{A['wd_share']}", USD)
line('vol_tr', 'Transferts entre portefeuilles', '$', lambda c, p, m: f"={c}{R['cli']}*{A['tr_volume']}", USD)
line('vol_total', 'Volume total qui passe par ScanLinkPay', '$', lambda c, p, m: f"={c}{R['vol_pay']}+{c}{R['vol_b2b']}+{c}{R['vol_api']}+{c}{R['vol_tr']}", USD, 'total')

psec('Recettes (prix payés, TVA comprise si applicable)')
line('r_comm', 'Commissions — paiements chez commerçants et entreprises', '$', lambda c, p, m: f"={c}{R['vol_pay']}*{A['comm']}", USD)
line('r_b2b', 'Commissions — commandes entre entreprises', '$', lambda c, p, m: f"={c}{R['vol_b2b']}*{A['comm']}", USD)
line('r_api', 'Commissions — applications partenaires (API)', '$',
     lambda c, p, m: '=' + '+'.join(f"{c}{R['api' + str(i - API_FIRST + 1)]}*'Hypothèses'!$F${i}" for i in range(API_FIRST, API_LAST + 1)), USD)
line('r_wd', 'Frais de retrait', '$', lambda c, p, m: f"={c}{R['vol_wd']}*{A['wd_fee']}", USD)
line('r_tr', 'Frais de transfert', '$', lambda c, p, m: f"={c}{R['vol_tr']}*{A['tr_fee']}", USD)
line('r_sub', 'Abonnements entreprises', '$', lambda c, p, m: f"={c}{R['e_pay']}*{A['sub_price']}", USD)
line('r_sup', 'Abonnements Fournisseur', '$', lambda c, p, m: f"={c}{R['sup_pay']}*{A['sup_price']}*((1-{A['sup_annual']})+{A['sup_annual']}*(12-{A['sup_free']})/12)", USD)
line('r_pro', 'Mes dépenses Pro', '$', lambda c, p, m: f"={c}{R['cli']}*{SC['pro']}*{A['pro_cdf']}/{A['fx']}", USD)
REV = ['r_comm', 'r_b2b', 'r_api', 'r_wd', 'r_tr', 'r_sub', 'r_sup', 'r_pro']
line('rev_ttc', 'Recettes encaissées', '$', lambda c, p, m: f"=SUM({c}{R['r_comm']}:{c}{R['r_pro']})", USD, 'total')
line('vat', 'TVA reversée à l’État', '$', lambda c, p, m: f"={c}{R['rev_ttc']}-{c}{R['rev_ttc']}/(1+{A['vat']}*{A['vat_on']})", USD)
line('rev', 'Chiffre d’affaires net (hors TVA)', '$', lambda c, p, m: f"={c}{R['rev_ttc']}-{c}{R['vat']}", USD, 'total')

psec('Coûts')
line('k_payout', 'FlexPaie — coût des retraits', '$', lambda c, p, m: f"={c}{R['vol_wd']}*{A['payout_cost']}", USD)
line('k_payin', 'FlexPaie — coût des encaissements (si à notre charge)', '$', lambda c, p, m: f"=({c}{R['vol_pay']}+{c}{R['vol_b2b']}+{c}{R['vol_api']})*{A['payin_cost']}", USD)
line('k_infra', 'Serveurs (selon le nombre d’utilisateurs)', '$', lambda c, p, m: f"={c}{R['users']}/1000*{A['infra_1000']}", USD)
line('k_sms', 'SMS', '$', lambda c, p, m: f"=({c}{R['e_new']}+{c}{R['m_new']}+{c}{R['c_new']})*{A['sms_new']}*{A['sms_price']}", USD)
line('k_acq', 'Acquisition (gagner de nouveaux utilisateurs)', '$', lambda c, p, m: f"={c}{R['e_new']}*{A['cac_e']}+{c}{R['m_new']}*{A['cac_m']}+{c}{R['c_new']}*{A['cac_c']}", USD)
line('k_fixed', 'Coûts fixes (équipe, hébergement, bureau…)', '$', lambda c, p, m: f"={A['fixed_total']}*(1+{A['team_growth']})^({c}$5-1)", USD)
line('cost', 'Total des coûts', '$', lambda c, p, m: f"=SUM({c}{R['k_payout']}:{c}{R['k_fixed']})", USD, 'total')

psec('Résultat')
line('ebitda', 'Résultat avant impôt', '$', lambda c, p, m: f"={c}{R['rev']}-{c}{R['cost']}", USD, 'total')
line('tax', 'Impôt sur les bénéfices (estimé)', '$', lambda c, p, m: f"=MAX(0,{c}{R['ebitda']})*{A['ibp']}", USD)
line('net', 'Bénéfice net', '$', lambda c, p, m: f"={c}{R['ebitda']}-{c}{R['tax']}", USD, 'total')
line('cum', 'Trésorerie cumulée (après investissement de départ)', '$', lambda c, p, m: (f"=-{A['init_total']}+{c}{R['net']}" if m == 1 else f"={p}{prow}+{c}{R['net']}"), USD, 'key')
line('margin', 'Marge nette', '%', lambda c, p, m: f"=IF({c}{R['rev']}=0,0,{c}{R['net']}/{c}{R['rev']})", PCT)
line('flag', 'Mois rentable (repère de calcul)', '', lambda c, p, m: f"=IF({c}{R['net']}>0,{c}$4,999)", NUM)
P.row_dimensions[R['flag']].hidden = True
line('flagcum', 'Trésorerie positive (repère de calcul)', '', lambda c, p, m: f"=IF({c}{R['cum']}>0,{c}$4,999)", NUM)
P.row_dimensions[R['flagcum']].hidden = True
line('flagobj', 'Objectif atteint (repère de calcul)', '', lambda c, p, m: f"=IF({c}{R['net']}>={A['target_gbp']}*{A['gbp']},{c}$4,999)", NUM)
P.row_dimensions[R['flagobj']].hidden = True
P.freeze_panes = 'C6'
for key in ('net', 'ebitda', 'cum'):
    rng = f'C{R[key]}:{LASTC}{R[key]}'
    P.conditional_formatting.add(rng, CellIsRule(operator='lessThan', formula=['0'], font=Font(name=F, color='C00000', bold=True)))
    P.conditional_formatting.add(rng, CellIsRule(operator='greaterThan', formula=['0'], font=Font(name=F, color='008000', bold=True)))

PR = lambda key: f"'Prévisions'!$C${R[key]}:${LASTC}${R[key]}"
YEAR = f"'Prévisions'!$C$5:${LASTC}$5"

# =====================================================================================================================
# 4. Synthèse (dashboard)
# =====================================================================================================================
Y = sheet('Synthèse', [50, 16, 16, 16, 16, 4, 44, 16])
title(Y, 'Synthèse sur 3 ans', None)
put(Y, 'A2', f'="Scénario : "&{A["scenario"]}&" · montants en USD (CDF au taux de "&TEXT({A["fx"]},"#,##0")&" pour 1 $)"', f=font(italic=True, color='555555'))
header(Y, 4, ['Indicateur', 'An 1', 'An 2', 'An 3', 'Total 3 ans'])
yrow = 5
YR = {}


def yline(key, label, kind='sum', fmt=USD, style=None):
    global yrow
    put(Y, f'A{yrow}', label, f=BOLD if style else font())
    for i, c_ in enumerate(('B', 'C', 'D'), 1):
        if kind == 'sum':
            f_ = f"=SUMIF({YEAR},{i},{PR(key)})"
        elif kind == 'end':  # value at the last month of the year
            f_ = f"=INDEX({PR(key)},{12 * i})"
        put(Y, f'{c_}{yrow}', f_, fmt, LINK)
    put(Y, f'E{yrow}', f'=SUM(B{yrow}:D{yrow})' if kind == 'sum' else f'=D{yrow}', fmt)
    if style:
        for c_ in 'ABCDE':
            Y[f'{c_}{yrow}'].fill = TOTAL
            Y[f'{c_}{yrow}'].font = BOLD
    YR[key] = yrow
    yrow += 1


def ysec(text):
    global yrow
    section(Y, yrow, text, 5)
    yrow += 1


ysec('Utilisateurs en fin d’année')
yline('e', 'Entreprises actives', 'end', NUM)
yline('mer', 'Commerçants actifs', 'end', NUM)
yline('cli', 'Clients actifs', 'end', NUM)
yline('sup_pay', 'Fournisseurs abonnés', 'end', NUM)
yrow += 1
ysec('Volumes')
yline('vol_total', 'Volume total qui passe par ScanLinkPay', 'sum', USD)
yline('vol_api', 'dont paiements via l’API des applications partenaires', 'sum', USD)
yrow += 1
ysec('Recettes par source (TVA comprise)')
labels = {
    'r_comm': 'Commissions — paiements commerçants / entreprises',
    'r_b2b': 'Commissions — commandes entre entreprises',
    'r_api': 'Commissions — applications partenaires (API)',
    'r_wd': 'Frais de retrait',
    'r_tr': 'Frais de transfert',
    'r_sub': 'Abonnements entreprises',
    'r_sup': 'Abonnements Fournisseur',
    'r_pro': 'Mes dépenses Pro',
}
for k in REV:
    yline(k, labels[k])
yline('rev', 'Chiffre d’affaires net (hors TVA)', style=True)
yrow += 1
ysec('Coûts et résultat')
yline('k_payout', 'FlexPaie — retraits')
yline('k_acq', 'Acquisition')
yline('k_fixed', 'Coûts fixes')
yline('cost', 'Total des coûts', style=True)
yline('ebitda', 'Résultat avant impôt', style=True)
yline('net', 'Bénéfice net', style=True)
put(Y, f'A{yrow}', 'Marge nette', bold=True)
for c_ in 'BCDE':
    put(Y, f'{c_}{yrow}', f'=IF({c_}{YR["rev"]}=0,0,{c_}{YR["net"]}/{c_}{YR["rev"]})', PCT)
YR['margin'] = yrow
yrow += 1
for key in ('ebitda', 'net'):
    rng = f'B{YR[key]}:E{YR[key]}'
    Y.conditional_formatting.add(rng, CellIsRule(operator='lessThan', formula=['0'], font=Font(name=F, color='C00000', bold=True)))

# KPI box on the right
put(Y, 'G4', 'Repères clés', f=WHITE_B, fill=HEAD)
put(Y, 'H4', '', fill=HEAD)
kpis = [
    ('Premier mois bénéficiaire', f"=IF(MIN('Prévisions'!$C${R['flag']}:${LASTC}${R['flag']})=999,\"Pas en 3 ans\",\"Mois \"&MIN('Prévisions'!$C${R['flag']}:${LASTC}${R['flag']}))", None),
    ('Investissement remboursé (trésorerie > 0)', f"=IF(MIN('Prévisions'!$C${R['flagcum']}:${LASTC}${R['flagcum']})=999,\"Pas en 3 ans\",\"Mois \"&MIN('Prévisions'!$C${R['flagcum']}:${LASTC}${R['flagcum']}))", None),
    ('Besoin de trésorerie maximal', f"=-MIN(0,MIN('Prévisions'!$C${R['cum']}:${LASTC}${R['cum']}))", USD),
    ('Bénéfice net du mois 36', f"=INDEX({PR('net')},36)", USD),
    ('Trésorerie cumulée à 3 ans', f"=INDEX({PR('cum')},36)", USD),
    ('Part des abonnements dans le CA (an 3)', f"=IF(D{YR['rev']}=0,0,(D{YR['r_sub']}+D{YR['r_sup']}+D{YR['r_pro']})/(1+{A['vat']}*{A['vat_on']})/D{YR['rev']})", PCT),
    ('Part de l’API partenaires dans le CA (an 3)', f"=IF(D{YR['rev']}=0,0,D{YR['r_api']}/(1+{A['vat']}*{A['vat_on']})/D{YR['rev']})", PCT),
]
for i, (lbl, f_, fm) in enumerate(kpis):
    rr = 5 + i
    put(Y, f'G{rr}', lbl)
    put(Y, f'H{rr}', f_, fm, BOLD).border = BOX

# Charts
# drop the unit column B from the series: build again precisely
ch = LineChart()
ch.title = 'Chiffre d’affaires, coûts et bénéfice net par mois (USD)'
ch.y_axis.title = 'USD'
ch.x_axis.title = 'Mois'
ch.height, ch.width = 8, 22
from openpyxl.chart import Series
for key, name in (('rev', 'Chiffre d’affaires net'), ('cost', 'Coûts'), ('net', 'Bénéfice net')):
    s = Series(Reference(P, min_col=FIRST, max_col=FIRST + MONTHS - 1, min_row=R[key], max_row=R[key]), title=name)
    ch.series.append(s)
ch.set_categories(Reference(P, min_col=FIRST, max_col=FIRST + MONTHS - 1, min_row=4, max_row=4))
Y.add_chart(ch, f'A{yrow + 2}')

bc = BarChart()
bc.type = 'col'
bc.grouping = 'stacked'
bc.overlap = 100
bc.title = 'D’où viennent les recettes (par an, USD)'
bc.height, bc.width = 8, 16
for k in REV:
    s = Series(Reference(Y, min_col=2, max_col=4, min_row=YR[k], max_row=YR[k]), title=labels[k])
    bc.series.append(s)
bc.set_categories(Reference(Y, min_col=2, max_col=4, min_row=4, max_row=4))
Y.add_chart(bc, f'G{yrow + 2}')

# =====================================================================================================================
# 5. Prix (subscriptions)
# =====================================================================================================================
X = sheet('Prix', [44, 14, 14, 14, 14, 14, 14, 14])
title(X, 'Fixer les prix des abonnements', 'Combien rapporte chaque prix, selon le nombre d’abonnés. Valeurs hors TVA (si TVA activée).')
put(X, 'A4', 'Diviseur TVA', f=font(size=9, color='777777'))
put(X, 'B4', f"=1+{A['vat']}*{A['vat_on']}", '0.00', LINK)
TV = "'Prix'!$B$4"

section(X, 6, 'A. Abonnement entreprise — recettes de l’an 1 selon le prix et la part qui paie', 8)
put(X, 'A7', 'Mois-entreprises actifs sur l’an 1 (scénario)', f=font(size=9))
put(X, 'B7', f"=SUMIF({YEAR},1,{PR('e')})", NUM, LINK)
header(X, 8, ['Prix / mois ($) ↓   Part qui paie →', 0.4, 0.5, 0.6, 0.7, 0.8])
for i in range(2, 7):
    X.cell(row=8, column=i).number_format = '0%'
prices = [5, 10, 15, 20, 25, 30]
for j, pz in enumerate(prices):
    rr = 9 + j
    put(X, f'A{rr}', pz, USD, INPUT).alignment = Alignment(horizontal='center')
    for i in range(2, 7):
        c_ = col(i)
        put(X, f'{c_}{rr}', f'=$A{rr}*{c_}$8*$B$7/{TV}', USD).border = BOX
X.conditional_formatting.add('B9:F14', FormulaRule(formula=[f"ABS($A9-{A['sub_price']})<0.01"], fill=PatternFill('solid', fgColor='FFF2CC')))
put(X, 'A15', 'Ligne surlignée = votre prix actuel. Plus le prix est haut, plus il faut de valeur visible (caisse, stock, rapports, fournisseurs) pour garder la même part qui paie.', f=font(size=9, italic=True, color='555555'))

section(X, 17, 'B. Abonnement Fournisseur — recettes par mois selon le prix et le nombre de fournisseurs abonnés', 8)
put(X, 'A18', 'Facteur annuel (part à l’année × mois offerts)', f=font(size=9))
put(X, 'B18', f"=(1-{A['sup_annual']})+{A['sup_annual']}*(12-{A['sup_free']})/12", '0.000', LINK)
header(X, 19, ['Prix / mois ($) ↓   Fournisseurs abonnés →', 10, 25, 50, 100, 200, 500])
for j, pz in enumerate([5, 8, 10, 15, 20, 25]):
    rr = 20 + j
    put(X, f'A{rr}', pz, USD, INPUT).alignment = Alignment(horizontal='center')
    for i in range(2, 8):
        c_ = col(i)
        put(X, f'{c_}{rr}', f'=$A{rr}*{c_}$19*$B$18/{TV}', USD).border = BOX
X.conditional_formatting.add('B20:G25', FormulaRule(formula=[f"ABS($A20-{A['sup_price']})<0.01"], fill=PatternFill('solid', fgColor='FFF2CC')))

section(X, 27, 'C. Formule annuelle Fournisseur — ce que coûte l’année selon les mois offerts', 8)
header(X, 28, ['Mois offerts', 'Prix de l’année ($)', 'Réduction', 'Équivalent par mois ($)'])
for j, n in enumerate(range(0, 5)):
    rr = 29 + j
    put(X, f'A{rr}', n, NUM, INPUT).alignment = Alignment(horizontal='center')
    put(X, f'B{rr}', f"={A['sup_price']}*(12-A{rr})", USD2).border = BOX
    put(X, f'C{rr}', f'=A{rr}/12', PCT).border = BOX
    put(X, f'D{rr}', f'=B{rr}/12', USD2).border = BOX

section(X, 35, 'D. Ce que rapporte un utilisateur par mois (hors TVA, scénario actif)', 8)
header(X, 36, ['Type d’utilisateur', 'Recette / mois ($)', 'Coût direct / mois ($)', 'Contribution ($)', 'Commentaire'])
X.column_dimensions['E'].width = 14
unit = [
    ('Commerçant', f"={SC['vol_m']}*({A['comm']}+{A['wd_share']}*{A['wd_fee']})/{TV}", f"={SC['vol_m']}*({A['wd_share']}*{A['payout_cost']}+{A['payin_cost']})", 'Commission + frais de retrait − coût FlexPaie.'),
    ('Entreprise (abonnée)', f"=({SC['vol_e']}*({A['comm']}+{A['wd_share']}*{A['wd_fee']})+{A['sub_price']})/{TV}", f"={SC['vol_e']}*({A['wd_share']}*{A['payout_cost']}+{A['payin_cost']})", 'Avec l’abonnement.'),
    ('Fournisseur (abonné)', f"=({A['sup_price']}*$B$18+{SC['ord']}*{SC['basket']}*{SC['ord_slp']}*({A['comm']}+{A['wd_share']}*{A['wd_fee']}))/{TV}", f"={SC['ord']}*{SC['basket']}*{SC['ord_slp']}*({A['wd_share']}*{A['payout_cost']}+{A['payin_cost']})", 'Abonnement + commissions sur ses ventes payées par ScanLinkPay.'),
    ('Client (avec Pro)', f"=({A['pro_cdf']}/{A['fx']}+{A['tr_volume']}*{A['tr_fee']})/{TV}", '=0', 'Le client paie peu : il fait vivre les commerçants.'),
]
for j, (lbl, rv, cv, note) in enumerate(unit):
    rr = 37 + j
    put(X, f'A{rr}', lbl)
    put(X, f'B{rr}', rv, USD2).border = BOX
    put(X, f'C{rr}', cv, USD2).border = BOX
    put(X, f'D{rr}', f'=B{rr}-C{rr}', USD2, BOLD).border = BOX
    put(X, f'E{rr}', note, f=font(size=9, color='555555'))
X.conditional_formatting.add('D37:D40', CellIsRule(operator='lessThan', formula=['0'], font=Font(name=F, color='C00000', bold=True)))

# =====================================================================================================================
# 6. Rentabilité — break-even and the £1 000 target; partner apps
# =====================================================================================================================
B = sheet('Rentabilité', [58, 18, 18, 18, 18, 18])
title(B, 'Seuil de rentabilité et objectif', 'Combien de transactions, de commerçants ou d’abonnés faut-il pour couvrir les coûts et atteindre l’objectif ?')
section(B, 4, 'A. Ce que rapporte 1 $ de paiement (hors TVA)', 6)
rows_b = [
    ('Commission par $ payé', f"={A['comm']}/'Prix'!$B$4", PCT2),
    ('Frais de retrait par $ payé (part retirée × frais)', f"={A['wd_share']}*{A['wd_fee']}/'Prix'!$B$4", PCT2),
    ('Coût FlexPaie par $ payé (retraits + encaissements)', f"={A['wd_share']}*{A['payout_cost']}+{A['payin_cost']}", PCT2),
    ('Contribution nette par $ payé', '=B5+B6-B7', PCT2),
    ('Contribution par transaction moyenne', f"=B8*{A['avg_pay']}", USD2),
]
for j, (lbl, f_, fm) in enumerate(rows_b):
    rr = 5 + j
    put(B, f'A{rr}', lbl, f=BOLD if j == 3 else font())
    put(B, f'B{rr}', f_, fm, BOLD if j == 3 else font()).border = BOX

section(B, 11, 'B. Couvrir les coûts fixes (an 1)', 6)
put(B, 'A12', 'Coûts fixes mensuels (an 1)')
put(B, 'B12', f"={A['fixed_total']}", USD, LINK).border = BOX
put(B, 'A13', 'Abonnements entreprises et Fournisseur (moyenne mensuelle an 1, hors TVA)')
put(B, 'B13', f"=(SUMIF({YEAR},1,{PR('r_sub')})+SUMIF({YEAR},1,{PR('r_sup')})+SUMIF({YEAR},1,{PR('r_pro')}))/12/'Prix'!$B$4", USD, LINK).border = BOX
put(B, 'A14', 'Reste à couvrir par les paiements')
put(B, 'B14', '=MAX(0,B12-B13)', USD).border = BOX
put(B, 'A15', 'Volume de paiements nécessaire par mois', bold=True)
put(B, 'B15', '=IF(B8<=0,"Impossible : la contribution par $ est négative",B14/B8)', USD, BOLD).border = BOX
put(B, 'A16', 'Soit, en transactions par mois')
put(B, 'B16', f"=IF(ISNUMBER(B15),B15/{A['avg_pay']},\"-\")", NUM).border = BOX
put(B, 'A17', 'Soit, en commerçants actifs')
put(B, 'B17', f"=IF(ISNUMBER(B15),B15/{SC['vol_m']},\"-\")", NUM1).border = BOX
put(B, 'A18', 'Ou, en entreprises abonnées seules (abonnement + paiements)')
put(B, 'B18', "=IF('Prix'!$D$38<=0,\"-\",B12/'Prix'!$D$38)", NUM1).border = BOX

section(B, 20, 'C. Objectif : bénéfice net de', 6)
B['A20'].value = f'="C. Objectif : bénéfice net de "&TEXT({A["target_gbp"]},"#,##0")&" £ par mois"'
put(B, 'A21', 'Objectif en $ (net d’impôt)')
put(B, 'B21', f"={A['target_gbp']}*{A['gbp']}", USD, LINK).border = BOX
put(B, 'A22', 'Résultat avant impôt nécessaire')
put(B, 'B22', f"=B21/(1-{A['ibp']})", USD).border = BOX
put(B, 'A23', 'Volume de paiements nécessaire par mois', bold=True)
put(B, 'B23', '=IF(B8<=0,"Impossible",(B22+B14)/B8)', USD, BOLD, YELLOW).border = BOX
put(B, 'A24', 'Soit, en transactions par mois')
put(B, 'B24', f"=IF(ISNUMBER(B23),B23/{A['avg_pay']},\"-\")", NUM, BOLD).border = BOX
put(B, 'A25', 'Soit, en transactions par jour')
put(B, 'B25', '=IF(ISNUMBER(B24),B24/30,"-")', NUM).border = BOX
put(B, 'A26', 'Soit, en commerçants actifs')
put(B, 'B26', f"=IF(ISNUMBER(B23),B23/{SC['vol_m']},\"-\")", NUM).border = BOX
put(B, 'A27', 'Mois où le modèle atteint l’objectif (scénario actif)')
put(B, 'B27', f"=IF(MIN('Prévisions'!$C${R['flagobj']}:${LASTC}${R['flagobj']})=999,\"Pas en 3 ans\",\"Mois \"&MIN('Prévisions'!$C${R['flagobj']}:${LASTC}${R['flagobj']}))", None, BOLD).border = BOX

section(B, 29, 'D. Applications partenaires (API) — combien rapporte chaque application', 6)
put(B, 'A30', 'Transactions par application et par mois', f=font())
put(B, 'B30', 500, NUM, INPUT, YELLOW).border = BOX
put(B, 'A31', 'Montant moyen par transaction ($)')
put(B, 'B31', 20, USD, INPUT, YELLOW).border = BOX
put(B, 'A32', 'Commission')
put(B, 'B32', f"={A['comm']}", PCT2, LINK).border = BOX
put(B, 'A33', 'Recette par application et par mois (hors TVA)', bold=True)
put(B, 'B33', "=B30*B31*B32/'Prix'!$B$4", USD2, BOLD).border = BOX
put(B, 'D30', 'Exemple HotelSaver : un hôtel paie son abonnement et ses clients leurs acomptes par l’API ScanLinkPay ; chaque paiement rapporte la commission.', f=font(size=9, italic=True, color='555555'))
B['D30'].alignment = Alignment(wrap_text=True, vertical='top')
B.merge_cells('D30:F33')
header(B, 35, ['Nombre d’applications', 'Recette / mois ($)', 'Recette / an ($)', 'Part des coûts fixes couverte'])
for j, n in enumerate([1, 2, 4, 6, 8, 10]):
    rr = 36 + j
    put(B, f'A{rr}', n, NUM, INPUT).alignment = Alignment(horizontal='center')
    put(B, f'B{rr}', f'=A{rr}*$B$33', USD).border = BOX
    put(B, f'C{rr}', f'=B{rr}*12', USD).border = BOX
    put(B, f'D{rr}', f'=IF($B$12=0,0,B{rr}/$B$12)', PCT).border = BOX
B.conditional_formatting.add('A36:D41', FormulaRule(formula=['$A36=4'], fill=PatternFill('solid', fgColor='FFF2CC')))
put(B, 'A43', 'Transactions nécessaires par application pour couvrir seule les coûts fixes', f=font())
put(B, 'B43', "=IF(B31*B32=0,\"-\",B12*'Prix'!$B$4/(B31*B32))", NUM, BOLD).border = BOX

# =====================================================================================================================
# 7. Retraits — fee vs FlexPaie cost
# =====================================================================================================================
W = sheet('Retraits', [40, 14, 14, 14, 14, 14, 14, 14])
title(W, 'Frais de retrait : gain ou perte ?', 'Le retrait vers Mobile Money coûte ~2 % chez FlexPaie. Nos frais doivent couvrir ce coût, plus la TVA.')
put(W, 'A4', 'Volume retiré par mois (moyenne an 1, scénario)')
put(W, 'B4', f"=SUMIF({YEAR},1,{PR('vol_wd')})/12", USD, LINK).border = BOX
put(W, 'A5', 'Coût FlexPaie par retrait')
put(W, 'B5', f"={A['payout_cost']}", PCT2, LINK).border = BOX
put(W, 'A6', 'Frais actuels (application)')
put(W, 'B6', f"={A['wd_fee']}", PCT2, LINK).border = BOX
header(W, 8, ['Frais de retrait', 'Recette hors TVA / 100 $', 'Coût / 100 $', 'Gain / 100 $', 'Gain par mois ($)', 'Verdict'])
for j, fee in enumerate([0.01, 0.015, 0.02, 0.025, 0.03, 0.035, 0.04]):
    rr = 9 + j
    put(W, f'A{rr}', fee, PCT, INPUT).alignment = Alignment(horizontal='center')
    put(W, f'B{rr}', f"=100*A{rr}/'Prix'!$B$4", USD2).border = BOX
    put(W, f'C{rr}', '=100*$B$5', USD2).border = BOX
    put(W, f'D{rr}', f'=B{rr}-C{rr}', USD2, BOLD).border = BOX
    put(W, f'E{rr}', f'=$B$4/100*D{rr}', USD).border = BOX
    put(W, f'F{rr}', f'=IF(D{rr}<0,"Perte",IF(D{rr}<0.3,"Juste","Rentable"))', None).border = BOX
W.conditional_formatting.add('D9:F15', FormulaRule(formula=['$D9<0'], font=Font(name=F, color='C00000', bold=True)))
W.conditional_formatting.add('D9:F15', FormulaRule(formula=['$D9>=0.3'], font=Font(name=F, color='008000', bold=True)))
W.conditional_formatting.add('A9:F15', FormulaRule(formula=[f'ABS($A9-$B$6)<0.0001'], fill=PatternFill('solid', fgColor='FFF2CC')))
put(W, 'A17', 'Recommandation : des frais de retrait d’au moins 2,5 % à 3 % (super admin → Frais et limites). La ligne surlignée = les frais actuels.', f=font(italic=True, color='555555'))

# =====================================================================================================================
# 8. Sensibilité — commission × volume
# =====================================================================================================================
Z = sheet('Sensibilité', [36, 14, 14, 14, 14, 14, 14])
title(Z, 'Sensibilité : commission et volume', 'Résultat avant impôt du mois 12 si l’on change la commission et le volume des paiements (le reste ne bouge pas).')
put(Z, 'A4', 'Résultat avant impôt du mois 12 (scénario)')
put(Z, 'B4', f"=INDEX({PR('ebitda')},12)", USD, LINK).border = BOX
put(Z, 'A5', 'Volume soumis à commission au mois 12')
put(Z, 'B5', f"=INDEX({PR('vol_pay')},12)+INDEX({PR('vol_b2b')},12)+INDEX({PR('vol_api')},12)", USD, LINK).border = BOX
put(Z, 'A6', 'Commission actuelle')
put(Z, 'B6', f"={A['comm']}", PCT2, LINK).border = BOX
put(Z, 'A7', 'Marge retraits par $ payé (hors TVA)')
put(Z, 'B7', f"={A['wd_share']}*({A['wd_fee']}/'Prix'!$B$4-{A['payout_cost']})", PCT2).border = BOX
header(Z, 9, ['Commission ↓   Volume →', 0.5, 0.75, 1, 1.5, 2, 3])
for i in range(2, 8):
    Z.cell(row=9, column=i).number_format = '0.00"x"'
for j, cm in enumerate([0.01, 0.015, 0.02, 0.025, 0.03, 0.035]):
    rr = 10 + j
    put(Z, f'A{rr}', cm, PCT, INPUT).alignment = Alignment(horizontal='center')
    for i in range(2, 8):
        c_ = col(i)
        f_ = f"=$B$4-$B$5*($B$6/'Prix'!$B$4+$B$7)+$B$5*{c_}$9*($A{rr}/'Prix'!$B$4+$B$7)"
        put(Z, f'{c_}{rr}', f_, USD).border = BOX
Z.conditional_formatting.add('B10:G15', CellIsRule(operator='lessThan', formula=['0'], font=Font(name=F, color='C00000', bold=True), fill=PatternFill('solid', fgColor='FDE2E2')))
Z.conditional_formatting.add('B10:G15', CellIsRule(operator='greaterThanOrEqual', formula=['0'], font=Font(name=F, color='006100'), fill=PatternFill('solid', fgColor='E2F5E6')))
put(Z, 'A17', 'Lecture : en rouge, le mois 12 est en perte. « 1x » = le volume prévu ; « 2x » = deux fois plus de paiements.', f=font(italic=True, color='555555'))

# =====================================================================================================================
# 0. Accueil (first tab)
# =====================================================================================================================
C = wb.create_sheet('Accueil', 0)
C.sheet_view.showGridLines = False
for i, w in enumerate([4, 42, 70], 1):
    C.column_dimensions[col(i)].width = w
for rr in range(1, 9):
    for cc in range(1, 4):
        C.cell(row=rr, column=cc).fill = PatternFill('solid', fgColor=NAVY)
put(C, 'B2', 'ScanLinkPay', f=font(size=26, bold=True, color='FFFFFF'), fill=HEAD)
put(C, 'B4', 'Modèle économique et étude de faisabilité', f=font(size=14, bold=True, color='FFFFFF'), fill=HEAD)
put(C, 'B5', 'Portefeuille mobile, paiements QR et carte, caisse et stock des entreprises, réseau de fournisseurs, API pour les applications partenaires', f=font(size=10, color='DDE3FF'), fill=HEAD)
put(C, 'B7', 'Version du 23 octobre 2026 — document de travail', f=font(size=9, italic=True, color='DDE3FF'), fill=HEAD)
put(C, 'B10', 'Résultat en un coup d’œil', f=font(size=12, bold=True, color=NAVY))
glance = [
    ('Scénario étudié', f"={A['scenario']}", None),
    ('Chiffre d’affaires net — an 1', f"='Synthèse'!B{YR['rev']}", USD),
    ('Chiffre d’affaires net — an 3', f"='Synthèse'!D{YR['rev']}", USD),
    ('Bénéfice net — total 3 ans', f"='Synthèse'!E{YR['net']}", USD),
    ('Premier mois bénéficiaire', "='Synthèse'!H5", None),
    ('Besoin de trésorerie maximal', "='Synthèse'!H7", USD),
    ('Paiements nécessaires pour l’objectif (par mois)', "='Rentabilité'!B23", USD),
]
for j, (lbl, f_, fm) in enumerate(glance):
    rr = 11 + j
    put(C, f'B{rr}', lbl).border = BOX
    cell = put(C, f'C{rr}', f_, fm, font(bold=True, color='008000'))
    cell.border = BOX
    cell.alignment = Alignment(horizontal='left')
put(C, 'B19', 'Les onglets', f=font(size=12, bold=True, color=NAVY))
tabs = [
    ('Hypothèses', 'Tous les prix, frais, coûts et applications partenaires. C’est ici qu’on modifie.'),
    ('Scénarios', 'Pessimiste / Réaliste / Optimiste : rythme d’inscription, volumes, départs.'),
    ('Prévisions', '36 mois : utilisateurs, volumes, recettes par source, coûts, bénéfice, trésorerie.'),
    ('Synthèse', 'Les 3 années côte à côte, repères clés et graphiques.'),
    ('Prix', 'Choisir les prix des abonnements (entreprise, Fournisseur, formule annuelle), recette par utilisateur.'),
    ('Rentabilité', 'Seuil de rentabilité, objectif de 1 000 £ par mois, ce que rapportent les applications partenaires.'),
    ('Retraits', 'Frais de retrait face au coût FlexPaie : gain ou perte.'),
    ('Sensibilité', 'Commission × volume : le résultat du mois 12.'),
]
for j, (t, d) in enumerate(tabs):
    rr = 20 + j
    put(C, f'B{rr}', t, f=font(bold=True, color=BRAND)).hyperlink = f"#'{t}'!A1"
    put(C, f'C{rr}', d).alignment = Alignment(wrap_text=True)
put(C, 'B29', 'Code couleur', f=font(size=12, bold=True, color=NAVY))
legend = [
    ('Texte bleu', 'Valeur saisie : hypothèse à discuter, modifiable', INPUT, None),
    ('Fond jaune', 'Hypothèse clé (elle change beaucoup le résultat)', font(), YELLOW),
    ('Texte vert', 'Valeur reprise d’un autre onglet', LINK, None),
    ('Texte noir', 'Calcul automatique : ne pas modifier', font(), None),
]
for j, (a, b, ft, fl) in enumerate(legend):
    rr = 30 + j
    put(C, f'B{rr}', a, f=ft, fill=fl).border = BOX
    put(C, f'C{rr}', b)
put(C, 'B35', 'Sources des tarifs', f=font(size=12, bold=True, color=NAVY))
src = [
    'Commission 2,5 % (min 0,05 $, max 100 $), retrait 1 %, transfert gratuit : valeurs par défaut du code ScanLinkPay (migrations 002, 007, 008), réglables par le super admin.',
    'FlexPaie : ~2 % sur les retraits (à notre charge), ~2,5 % à 3 % sur les encaissements (payés par le client) — docs/FLEXPAIE.md.',
    'Abonnement Fournisseur 10 $ / mois, mois offerts sur l’année et essai réglés par le super admin (décision du 23/10/2026, migration 084).',
    'Mes dépenses Pro 5 000 CDF / mois (migration 025). Prix de l’abonnement entreprise : à remplacer par le prix réel.',
    'Volumes, adoption, coûts d’équipe et de marketing : hypothèses de travail à discuter, pas des chiffres constatés.',
]
for j, t in enumerate(src):
    rr = 36 + j
    put(C, f'B{rr}', '•')
    put(C, f'C{rr}', t, f=font(size=9)).alignment = Alignment(wrap_text=True, vertical='top')
    C.row_dimensions[rr].height = 26

for ws in wb.worksheets:
    ws.sheet_properties.tabColor = NAVY if ws.title in ('Accueil', 'Synthèse') else BRAND if ws.title in ('Hypothèses', 'Scénarios') else '8EA2F5'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
wb.active = 0
wb.save(OUT)
print('saved', OUT, R['net'], R['rev'])
